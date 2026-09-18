"""Strict, staged conversion of uploaded official templates (no uploaded Jinja)."""
from __future__ import annotations

import io
from pathlib import Path, PurePosixPath
import tempfile
import zipfile

from docx import Document
from docxtpl import DocxTemplate
from lxml import etree
from werkzeug.utils import secure_filename

from convert_template import convert_template
from generator import _build_context, _render_one
from template_storage import LANGUAGES, activate_templates, template_lock

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_REQUEST_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
NAMESPACE = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


class TemplateUpdateError(ValueError):
    """A conversion error safe to show to the uploader."""


def validate_docx(path: Path, *, official: bool = False) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise TemplateUpdateError("DOCX expanded contents exceed the safety limit")
            if len(names) != len(set(names)):
                raise TemplateUpdateError("DOCX contains duplicate ZIP entries")
            for name in names:
                if "\\" in name or PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts:
                    raise TemplateUpdateError("DOCX contains unsafe ZIP paths")
                lowered = name.lower()
                if any(token in lowered for token in ("vbaproject", "embeddings/", "activex/", "customui/")):
                    raise TemplateUpdateError("Macro or embedded executable content is not accepted")
            for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
                if required not in names:
                    raise TemplateUpdateError(f"DOCX is missing {required}")
            if archive.testzip():
                raise TemplateUpdateError("DOCX ZIP integrity check failed")
            for name in names:
                if not name.endswith((".xml", ".rels")):
                    continue
                data = archive.read(name)
                if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                    raise TemplateUpdateError("DOCX XML entity declarations are not accepted")
                root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                if name == "[Content_Types].xml":
                    content = data.lower()
                    if b"macroenabled" in content or b"vba" in content or b"oleobject" in content:
                        raise TemplateUpdateError("Macro-enabled or embedded object documents are not accepted")
                    if b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml" not in content:
                        raise TemplateUpdateError("Please upload a standard DOCX Word document")
                if name.endswith(".rels"):
                    for rel in root:
                        if rel.get("TargetMode") == "External" and (
                            not rel.get("Type", "").endswith("/hyperlink")
                            or not rel.get("Target", "").lower().startswith(("https://", "http://", "mailto:"))
                        ):
                            raise TemplateUpdateError("External linked content other than web/email hyperlinks is not accepted")
                if official:
                    text = etree.tostring(root, encoding="unicode")
                    text += "".join(root.itertext()) + "".join(root.xpath("//w:t/text()", namespaces=NAMESPACE))
                    text += "".join(value for node in root.iter() for value in node.attrib.values())
                    if any(token in text for token in ("{{", "}}", "{%", "%}", "{#", "#}", "__MARKING_RULE_BLOCK__")):
                        raise TemplateUpdateError("Upload an official template without runtime/Jinja placeholders")
                    if root.xpath("//*[local-name()='altChunk' or local-name()='object']"):
                        raise TemplateUpdateError("Embedded or imported content is not accepted")
        Document(path)
    except TemplateUpdateError:
        raise
    except Exception as exc:
        raise TemplateUpdateError("Malformed or unreadable DOCX document") from exc


def save_upload(upload, destination: Path, language: str) -> None:
    filename = upload.filename or ""
    # Filename is never used as a path; rejecting traversal also gives useful feedback.
    if "/" in filename or "\\" in filename or ".." in filename or not secure_filename(filename):
        raise TemplateUpdateError(f"{language.upper()}: unsafe upload filename")
    extension = Path(filename).suffix.lower()
    if extension == ".doc":
        raise TemplateUpdateError(f"{language.upper()}: export the .doc file to .docx using Microsoft Word or Apple Pages first")
    if extension != ".docx":
        raise TemplateUpdateError(f"{language.upper()}: only .docx files are accepted (not .docm)")
    payload = upload.stream.read(MAX_FILE_BYTES + 1)
    if not payload or len(payload) > MAX_FILE_BYTES:
        raise TemplateUpdateError(f"{language.upper()}: file must be non-empty and at most 10 MB")
    destination.write_bytes(payload)
    validate_docx(destination, official=True)


def validate_prepared(path: Path, language: str) -> None:
    validate_docx(path)
    context = _build_context({}, language)
    variables = DocxTemplate(str(path)).get_undeclared_template_variables()
    required = set(context) - {"degree_level"}  # Attendance contains the degree; no separate slot.
    if variables != required:
        raise TemplateUpdateError(
            f"{language.upper()}: incorrect placeholders; missing={sorted(required - variables)}, "
            f"unexpected={sorted(variables - required)}"
        )
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        text = "".join(root.itertext())
        if text.count("__MARKING_RULE_BLOCK__") != 1:
            raise TemplateUpdateError(f"{language.upper()}: expected exactly one marking-rule slot")
    # Exercise every rule, degree and missing-value handling, including joint content.
    for rule in (1, 2, 3, 4):
        for degree in ("bachelor", "master", "doctoral"):
            sample = {
                "class_code": "TEST1001-111, TEST1001-114",
                "class_codes": ["TEST1001-111", "TEST1001-114"],
                "prog_name_en": "Programme A / Programme B",
                "module_name_en": "Validation & <Module>",
                "faculty_en": "Test Faculty", "rule_code": rule,
                "degree_level": degree, "academic_year": "2026/2027", "semester": "1",
                "medium_of_instruction": None,
            }
            rendered = _render_one(sample, language, template_path=path)
            doc = Document(io.BytesIO(rendered))
            if doc.tables[0].cell(1, 1).text != sample["class_code"] or doc.tables[0].cell(4, 1).text:
                raise TemplateUpdateError(f"{language.upper()}: representative metadata generation failed")
            if "Programme A / Programme B" not in " ".join(p.text for p in doc.paragraphs):
                raise TemplateUpdateError(f"{language.upper()}: representative programme generation failed")


def update_templates(uploads: dict, active_dir: Path) -> str:
    missing = [lang.upper() for lang in LANGUAGES if lang not in uploads or not uploads[lang].filename]
    if missing:
        raise TemplateUpdateError(f"Upload all three official DOCX templates. Missing: {', '.join(missing)}")
    active_dir = Path(active_dir)
    # Serialize the whole upload lifecycle so orphan cleanup after a crash can
    # never delete another worker's live staging directory.
    with template_lock(active_dir), tempfile.TemporaryDirectory(prefix=".template-stage-", dir=active_dir.parent) as temp:
        staging = Path(temp)
        for lang in LANGUAGES:
            source = staging / f"official_{lang}.docx"
            prepared = staging / f"template_{lang}.docx"
            try:
                save_upload(uploads[lang], source, lang)
                convert_template(source, prepared, lang)
                validate_prepared(prepared, lang)
            except Exception as exc:
                raise TemplateUpdateError(f"{lang.upper()} template validation failed: {exc}") from exc
        return activate_templates(staging, active_dir)
