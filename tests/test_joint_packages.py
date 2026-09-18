"""Exercise the real Flask selection, SQLite resolution and ZIP generation path."""
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document

from test_generator import ROOT
import app
import database
import generator


class JointPackageRouteTests(unittest.TestCase):
    codes = ["COMP111-111", "COMP1121-111", "COMP1121-114"]

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db = self.root / "classes.db"
        database.init_db(self.db, seed=False)
        self.db_patch = patch.object(database, "DB_PATH", self.db)
        self.db_patch.start()
        self.addCleanup(self.db_patch.stop)
        self.output_patch = patch.object(generator, "OUTPUT_DIR", self.root / "output")
        self.output_patch.start()
        self.addCleanup(self.output_patch.stop)
        self.ready_patch = patch.object(app, "DATABASE_STARTUP_ERROR", None)
        self.ready_patch.start()
        self.addCleanup(self.ready_patch.stop)
        with database.get_connection() as conn:
            conn.execute("INSERT INTO faculties (id, code) VALUES (1, 'FCA')")
            for index in (1, 2, 3):
                conn.execute(
                    "INSERT INTO programmes (id, code, name_en, name_zh, name_pt, degree_level, faculty_id) VALUES (?, ?, ?, ?, ?, 'bachelor', 1)",
                    (index, f"P{index}", f"Programme {index}", f"課程{index}", f"Curso {index}"),
                )
            for index, code in enumerate(self.codes + ["SOLO1-111", "SOLO2-111"], 1):
                conn.execute(
                    "INSERT INTO classes (id, class_code, module_code, module_name_en, rule_code, programme_id, joint_relationship) VALUES (?, ?, ?, 'Module', 2, ?, ?)",
                    (index, code, code.split('-')[0], min(index, 3), ', '.join(c for c in self.codes if c != code) if index <= 3 else ''),
                )
        self.client = app.app.test_client()

    def check_download(self, payload, codes, joint_codes=None):
        payload = dict(payload, academic_year="2026/2027", semester="1")
        with self.client.post("/api/generate", json=payload) as response:
            self.assertEqual(200, response.status_code, response.get_json(silent=True))
            self.assertIn("module_outlines_2026_2027_sem1.zip", response.headers["Content-Disposition"])
            content = response.data
        count = 0
        with zipfile.ZipFile(io.BytesIO(content)) as outer:
            self.assertEqual(sorted(f"{code}.zip" for code in codes), outer.namelist())
            self.assertIsNone(outer.testzip())
            for code in codes:
                with zipfile.ZipFile(io.BytesIO(outer.read(f"{code}.zip"))) as inner:
                    self.assertEqual([f"{code}_{lang}.docx" for lang in ("EN", "ZH", "PT")], inner.namelist())
                    self.assertIsNone(inner.testzip())
                    for lang in ("EN", "ZH", "PT"):
                        doc = Document(io.BytesIO(inner.read(f"{code}_{lang}.docx")))
                        expected_codes = (joint_codes or self.codes) if code in self.codes else [code]
                        self.assertEqual(", ".join(expected_codes), doc.tables[0].cell(1, 1).text)
                        if code in self.codes:
                            prefix = {"EN": "Programme ", "ZH": "課程", "PT": "Curso "}[lang]
                            programmes = " / ".join(f"{prefix}{self.codes.index(c) + 1}" for c in expected_codes)
                            self.assertEqual(programmes, doc.paragraphs[1].text.strip())
                        count += 1
        self.assertEqual(len(codes) * 3, count)

    def test_three_joint_members_have_three_packages_and_nine_joint_docs(self):
        self.check_download({"class_ids": [1, 2, 3]}, self.codes)

    def test_two_joint_members_have_two_packages_and_six_joint_docs(self):
        with database.get_connection() as conn:
            conn.execute("DELETE FROM classes WHERE id = 3")
        self.check_download({"class_ids": [1, 2]}, self.codes[:2], self.codes[:2])

    def test_standalone_has_one_package_and_three_docs(self):
        self.check_download({"class_ids": [4]}, ["SOLO1-111"])

    def test_faculty_mixed_batch_has_five_packages_and_fifteen_docs(self):
        self.check_download({"faculty_id": 1}, self.codes + ["SOLO1-111", "SOLO2-111"])

    def test_selecting_one_joint_member_only_outputs_that_member(self):
        for member_id, code in enumerate(self.codes, 1):
            with self.subTest(code=code):
                self.check_download({"class_ids": [member_id]}, [code])

    def test_multiple_selected_members_are_not_deduplicated_by_group(self):
        self.check_download({"class_ids": [2, 3, 2]}, self.codes[1:])

    def test_programme_only_outputs_its_members_with_complete_joint_content(self):
        self.check_download({"programme_id": 1}, [self.codes[0]])

    def test_faculty_scope_does_not_output_joint_members_in_other_faculties(self):
        with database.get_connection() as conn:
            conn.execute("INSERT INTO faculties (id, code) VALUES (2, 'OTHER')")
            conn.execute("UPDATE programmes SET faculty_id = 2 WHERE id = 1")
        self.check_download({"faculty_id": 2}, [self.codes[0]])

    def test_classes_api_exposes_actual_member_ids_and_full_joint_metadata(self):
        choices = self.client.get("/api/classes?faculty_id=1").get_json()
        self.assertEqual(5, len(choices))
        for member_id, code in enumerate(self.codes, 1):
            choice = next(c for c in choices if c["id"] == member_id)
            self.assertEqual(code, choice["class_code"])
            self.assertEqual(self.codes, choice["class_codes"])
            self.assertEqual([code], choice["output_class_codes"])
        programme_choices = self.client.get("/api/classes?programme_id=1").get_json()
        self.assertEqual([self.codes[0]], [c["class_code"] for c in programme_choices])

    def test_unselected_joint_member_rule_conflict_blocks_generation(self):
        with database.get_connection() as conn:
            conn.execute("UPDATE classes SET rule_code = 4 WHERE id = 3")
        for payload in ({"class_ids": [1]}, {"programme_id": 1}, {"faculty_id": 1}):
            with self.subTest(payload=payload):
                with self.client.post("/api/generate", json=payload) as response:
                    self.assertEqual(400, response.status_code)
                    self.assertIn("conflicting Rule", response.get_json()["error"])
                    for code in self.codes:
                        self.assertIn(code, response.get_json()["error"])
        self.assertFalse((self.root / "output").exists())
