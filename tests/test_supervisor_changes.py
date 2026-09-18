"""End-to-end regressions for per-class packages and safe template administration."""
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from werkzeug.datastructures import FileStorage

from test_generator import ROOT, sample_class, word_xml
import app
import generator
import template_storage
import template_update
from rules import RULE_PARAGRAPHS, JointRuleConflictError


def official_copy(language):
    """Restore official slots in a known-good runtime copy, retaining formatting.

    The local PT .docx source is a flattened legacy export, so use the working
    5-column runtime as the reference for its official-slot fixture.
    """
    headers = {
        "en": ("[Name of academic unit]", "[Programme name]", "[Insert marking scheme]"),
        "zh": ("[學術單位名稱]", "[課程名稱]", "[插入評分準則]"),
        "pt": ("[nome da unidade académica]", "[designação do curso]", "[Inserir o critério de classificação]"),
    }
    import re
    output = io.BytesIO()
    with zipfile.ZipFile(generator.TEMPLATES[language]) as source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.endswith(".xml"):
                text = content.decode()
                unit, programme, marking = headers[language]
                text = text.replace("{{ academic_unit }}", unit).replace("{{ programme_name }}", programme)
                text = text.replace("{{ attendance_text }}", generator._attendance_text("bachelor", language))
                text = text.replace(generator.MARKING_RULE_MARKER, marking)
                text = re.sub(r"{{\s*\w+\s*}}", "", text)
                content = text.encode()
            target.writestr(info, content)
    return output.getvalue()


def mutate_docx(payload, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            target.writestr(info, transform(info.filename, content))
    return output.getvalue()


class PackageTests(unittest.TestCase):
    def check_packages(self, classes, expected_codes):
        with tempfile.TemporaryDirectory() as output:
            buffer = generator.generate_batch(classes, output_dir=output)
            with zipfile.ZipFile(buffer) as outer:
                self.assertEqual(sorted(f"{code}.zip" for code in expected_codes), outer.namelist())
                self.assertIsNone(outer.testzip())
                for code in expected_codes:
                    with zipfile.ZipFile(io.BytesIO(outer.read(f"{code}.zip"))) as inner:
                        self.assertEqual([f"{code}_{lang}.docx" for lang in ("EN", "ZH", "PT")], inner.namelist())
                        self.assertIsNone(inner.testzip())
                        group = next(c for c in classes if code in (c.get("class_codes") or [c["class_code"]]))
                        for lang in ("en", "zh", "pt"):
                            document = Document(io.BytesIO(inner.read(f"{code}_{lang.upper()}.docx")))
                            self.assertEqual(", ".join(group.get("class_codes") or [code]), document.tables[0].cell(1, 1).text)
                            programme = group.get(f"prog_name_{lang}") or group["prog_name_en"]
                            self.assertEqual(programme, document.paragraphs[1].text.strip())
            self.assertEqual(len(expected_codes) * 3, len(list(Path(output).glob("generated_*/*.docx"))))
        return buffer.getvalue()

    def test_single_standalone(self):
        self.check_packages([sample_class(class_code="COMP1121-111")], ["COMP1121-111"])

    def test_two_standalone(self):
        self.check_packages([sample_class(class_code=c) for c in ("A-111", "B-111")], ["A-111", "B-111"])

    def test_two_and_three_member_joint_groups(self):
        for codes in (["A-111", "A-114"], ["COMP111-111", "COMP1121-111", "COMP1121-114"]):
            with self.subTest(codes=codes):
                self.check_packages([sample_class(class_code=", ".join(codes), class_codes=codes,
                    prog_name_en="Programme One / Programme Two", prog_name_zh="課程一 / 課程二", prog_name_pt="Curso Um / Curso Dois")], codes)

    def test_batch_two_standalone_plus_three_joint_members(self):
        joint = sample_class(class_code="J1, J2, J3", class_codes=["J1", "J2", "J3"])
        self.check_packages([sample_class(class_code="A"), sample_class(class_code="B"), joint], ["A", "B", "J1", "J2", "J3"])

    def test_repeated_group_does_not_duplicate_packages(self):
        joint = sample_class(class_code="A, B", class_codes=["A", "B"])
        self.check_packages([joint, joint], ["A", "B"])

    def test_filename_collisions_are_reported(self):
        with tempfile.TemporaryDirectory() as output:
            with self.assertRaisesRegex(ValueError, "same package filename"):
                generator.generate_batch([sample_class(class_code="A/B"), sample_class(class_code="A_B")], output_dir=output)
            self.assertEqual([], list(Path(output).iterdir()))

    def test_conflict_blocks_whole_batch_before_output(self):
        joint = sample_class(class_code="A, B", class_codes=["A", "B"], rule_codes=[2, 4])
        with tempfile.TemporaryDirectory() as output:
            with self.assertRaises(JointRuleConflictError):
                generator.generate_batch([sample_class(class_code="C"), joint], output_dir=output)
            self.assertEqual([], list(Path(output).iterdir()))

    def test_faculty_route_generates_all_class_packages(self):
        classes = [sample_class(class_code="A"), sample_class(class_code="B, C", class_codes=["B", "C"])]
        with tempfile.TemporaryDirectory() as output, patch.object(app, "get_classes_full", return_value=classes) as fetch:
            real_generate = generator.generate_batch
            with patch.object(app, "generate_batch", side_effect=lambda rows, **kw: real_generate(rows, output_dir=output, **kw)):
                response = app.app.test_client().post("/api/generate", json={"faculty_id": 1, "academic_year": "2026/2027", "semester": "1"})
        fetch.assert_called_once_with(faculty_id=1)
        self.assertEqual(200, response.status_code)
        self.assertIn("module_outlines_2026_2027_sem1.zip", response.headers["Content-Disposition"])
        with zipfile.ZipFile(io.BytesIO(response.data)) as outer:
            self.assertEqual(["A.zip", "B.zip", "C.zip"], outer.namelist())
        response.close()


class RuleReferenceTests(unittest.TestCase):
    def test_api_uses_canonical_wording(self):
        response = app.app.test_client().get("/api/rules")
        self.assertEqual(200, response.status_code)
        rows = response.get_json()
        self.assertEqual([1, 2, 3, 4], [r["code"] for r in rows])
        for row in rows:
            expected = RULE_PARAGRAPHS[row["code"]]
            self.assertEqual({lang: list(text) for lang, text in expected.items()}, row["paragraphs"])
            self.assertEqual(" ".join(expected["en"]) or "No additional marking-rule condition.", row["summary"])

    def test_rules_available_during_database_migration_review(self):
        with patch.object(app, "DATABASE_STARTUP_ERROR", "Review required"):
            self.assertEqual(200, app.app.test_client().get("/api/rules").status_code)

    def test_frontend_rules_render_and_actual_package_counts(self):
        # Run the actual frontend helpers with a small DOM stub, no duplicate implementation.
        import re
        html = (ROOT / "frontend/index.html").read_text()
        helpers = re.search(r"  function packageCount\(.*?(?=  \(async \(\) =>)", html, re.S).group()
        script = '''const assert = require('node:assert/strict');
const document = {
 createElement: tag => ({tag, children: [], textContent: '', appendChild(child) { this.children.push(child); }}),
 createTextNode: text => ({textContent: text}),
};
'''+helpers+'''\nconst box = {children: [], replaceChildren() {this.children=[];}, appendChild(x) {this.children.push(x);}};
renderRules([{code:1, summary:'Canonical reminder'}, {code:2, summary:'<script>literal</script>'}], box);
assert.equal(box.children.length, 2);
assert.equal(box.children[0].children[0].textContent, 'Rule 1 — ');
assert.equal(box.children[1].children[1].textContent, '<script>literal</script>');
assert.equal(packageCount([{class_codes:['A']}, {class_codes:['B']}, {class_codes:['J1','J2','J3']}]), 5);
assert.equal(packageCount([{class_codes:['J1','J2','J3']}]), 3);
assert.equal(packageCount([{class_codes:['A','B']}, {class_codes:['A','B']}]), 2);
assert.equal(packageCount([{class_code:'A', class_codes:['A','B','C'], output_class_codes:['A']}]), 1);
assert.equal(packageCount([{class_codes:['A','B','C'], output_class_codes:['A']}, {class_codes:['A','B','C'], output_class_codes:['B']}]), 2);
'''
        result = subprocess.run(["node", "-e", script], capture_output=True, text=True)
        self.assertEqual(0, result.returncode, result.stderr)
        self.assertIn('fetch("/api/rules")', html)
        self.assertNotIn("Joint classes count as one outline", html)
        self.assertIn('fetch("/api/templates/convert"', html)
        self.assertIn('res.headers.get("Content-Disposition")', html)


class TemplateUpdateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sources = {lang: official_copy(lang) for lang in ("en", "zh", "pt")}

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.active = Path(self.temp.name).resolve() / "templates"
        self.active.mkdir()
        for lang in ("en", "zh", "pt"):
            shutil.copyfile(generator.TEMPLATES[lang], self.active / f"template_{lang}.docx")
        self.before = self.snapshot()
        self.config = patch.dict(app.app.config, TEMPLATE_ADMIN_TOKEN="test-admin-token", TEMPLATE_DIR=self.active, TEMPLATE_STORAGE_CONFIGURED=True)
        self.config.start()
        self.addCleanup(self.config.stop)
        self.client = app.app.test_client()

    def snapshot(self):
        return {p.name: p.read_bytes() for p in self.active.glob("*.docx")}

    def post(self, changes=None, missing=(), authorization="Bearer test-admin-token"):
        files = {lang: (io.BytesIO(content), f"official_{lang}.docx") for lang, content in self.sources.items() if lang not in missing}
        files.update(changes or {})
        response = self.client.post("/api/templates/convert", data=files, headers={"Authorization": authorization})
        # The test client owns the spooled WSGI input (normally the web server's
        # responsibility), including bodies rejected before multipart parsing.
        response.request.environ["wsgi.input"].close()
        return response

    def assert_preserved(self, response, status=400):
        self.assertEqual(status, response.status_code, response.get_json())
        self.assertIn("error", response.get_json())
        self.assertEqual(self.before, self.snapshot())

    def test_valid_conversion_activates_all_backups_and_representative_generation(self):
        response = self.post()
        self.assertEqual(200, response.status_code, response.get_json())
        backup = self.active.parent / ".template-backups" / response.get_json()["backup_id"]
        self.assertEqual(self.before, {p.name: p.read_bytes() for p in backup.glob("*.docx")})
        for lang in ("en", "zh", "pt"):
            path = self.active / f"template_{lang}.docx"
            self.assertNotEqual(self.before[path.name], path.read_bytes())
            for rule in (1, 2, 3, 4):
                rendered = generator._render_one(sample_class(rule_code=rule), lang, template_path=path)
                generator._validate_rendered_docx(rendered)
                self.assertEqual(6, len(Document(io.BytesIO(rendered)).tables))
                baseline = generator._render_one(sample_class(rule_code=rule), lang)
                with zipfile.ZipFile(io.BytesIO(baseline)) as old, zipfile.ZipFile(io.BytesIO(rendered)) as new:
                    self.assertEqual(set(old.namelist()), set(new.namelist()))
                    for name in old.namelist():
                        self.assertEqual(old.read(name), new.read(name), f"Conversion changed formatting/content: {lang}/{name}")
                if lang in ("en", "zh"):
                    self.assertIn("https://mpusite.mpu.edu.mo/studenthandbook/", word_xml(rendered))
                    self.assertNotIn("teaching_learning/", word_xml(rendered))
        self.assertFalse((self.active.parent / ".template-transaction.json").exists())
        self.assertFalse(list(self.active.parent.glob(".template-stage-*")))

    def test_missing_languages(self):
        for lang in ("en", "zh", "pt"):
            with self.subTest(lang=lang):
                response = self.post(missing=(lang,))
                self.assert_preserved(response)
                self.assertIn(lang.upper(), response.get_json()["error"])

    def test_invalid_extensions_and_portuguese_doc_guidance(self):
        for extension in (".txt", ".docm", ".doc", ".zip"):
            with self.subTest(extension=extension):
                response = self.post({"pt": (io.BytesIO(self.sources["pt"]), "official" + extension)})
                self.assert_preserved(response)
                if extension == ".doc":
                    self.assertIn("Microsoft Word or Apple Pages", response.get_json()["error"])

    def test_malformed_docx_and_broken_xml(self):
        self.assert_preserved(self.post({"en": (io.BytesIO(b"not a zip"), "official.docx")}))
        broken = mutate_docx(self.sources["en"], lambda name, data: b"<broken" if name == "word/document.xml" else data)
        self.assert_preserved(self.post({"en": (io.BytesIO(broken), "official.docx")}))

    def test_wrong_structure_and_language_do_not_activate(self):
        empty = io.BytesIO()
        Document().save(empty)
        self.assert_preserved(self.post({"pt": (io.BytesIO(empty.getvalue()), "official.docx")}))
        self.assert_preserved(self.post({"zh": (io.BytesIO(self.sources["en"]), "official.docx")}))

    def test_conversion_failure_preserves_all_templates(self):
        real = template_update.convert_template
        def fail(src, dst, lang):
            if lang == "pt":
                raise ValueError("simulated Portuguese failure")
            return real(src, dst, lang)
        with patch.object(template_update, "convert_template", side_effect=fail):
            self.assert_preserved(self.post())
        self.assertFalse((self.active.parent / ".template-backups").exists())

    def test_activation_failure_rolls_back_all_files(self):
        real = template_storage._atomic_copy
        failed = False
        def fail_once(src, dst):
            nonlocal failed
            if dst == self.active / "template_zh.docx" and ".template-stage-" in str(src) and not failed:
                failed = True
                raise OSError("simulated disk failure")
            return real(src, dst)
        with patch.object(template_storage, "_atomic_copy", side_effect=fail_once), self.assertLogs(app.app.logger, level="ERROR"):
            self.assert_preserved(self.post(), status=500)
        self.assertTrue(failed)
        self.assertFalse((self.active.parent / ".template-transaction.json").exists())
        self.assertEqual(1, len(list((self.active.parent / ".template-backups").iterdir())))

    def test_traversal_filenames_rejected(self):
        for name in ("../../template_en.docx", "..\\template_en.docx", "/tmp/template.docx"):
            with self.subTest(name=name):
                self.assert_preserved(self.post({"en": (io.BytesIO(self.sources["en"]), name)}))

    def test_readable_filename_with_spaces_is_safe(self):
        upload = FileStorage(stream=io.BytesIO(self.sources["en"]), filename="English Official Template.docx")
        dest = self.active.parent / "canonical.docx"
        template_update.save_upload(upload, dest, "en")
        self.assertEqual(self.sources["en"], dest.read_bytes())

    def test_upload_jinja_rejected_without_execution(self):
        bad = mutate_docx(self.sources["en"], lambda name, data: data.replace(b"[Name of academic unit]", b"{{ unsafe.attribute }}") if name == "word/document.xml" else data)
        response = self.post({"en": (io.BytesIO(bad), "official.docx")})
        self.assert_preserved(response)
        self.assertIn("without runtime/Jinja", response.get_json()["error"])

    def test_upload_jinja_in_xml_comment_rejected(self):
        bad = mutate_docx(self.sources["en"], lambda name, data: data.replace(b"<w:body>", b"<w:body><!-- {{ unsafe.attribute }} -->") if name == "word/document.xml" else data)
        response = self.post({"en": (io.BytesIO(bad), "official.docx")})
        self.assert_preserved(response)
        self.assertIn("without runtime/Jinja", response.get_json()["error"])

    def test_unresolved_placeholder_after_conversion_blocks_activation(self):
        real = template_update.convert_template
        def inject(src, dst, lang):
            real(src, dst, lang)
            if lang == "pt":
                content = mutate_docx(dst.read_bytes(), lambda name, data: data.replace(b"{{ module_name }}", b"{{ unknown_slot }}") if name == "word/document.xml" else data)
                dst.write_bytes(content)
        with patch.object(template_update, "convert_template", side_effect=inject):
            response = self.post()
        self.assert_preserved(response)
        self.assertIn("incorrect placeholders", response.get_json()["error"])

    def test_size_limit_and_empty_file(self):
        for data in (b"", b"x" * (template_update.MAX_FILE_BYTES + 1)):
            self.assert_preserved(self.post({"en": (io.BytesIO(data), "official.docx")}))

    def test_total_request_limit(self):
        with patch.object(app, "MAX_REQUEST_BYTES", 100):
            self.assert_preserved(self.post(), status=413)

    def test_auth_required_and_disabled_by_default(self):
        self.assert_preserved(self.post(authorization=""), status=401)
        self.assert_preserved(self.post(authorization="Bearer incorrect"), status=401)
        with patch.dict(app.app.config, TEMPLATE_ADMIN_TOKEN=""):
            self.assert_preserved(self.post(), status=503)

    def test_storage_configuration_required_before_upload_processing(self):
        with patch.dict(app.app.config, TEMPLATE_STORAGE_CONFIGURED=False), patch.object(app, "update_templates") as update:
            self.assert_preserved(self.post(), status=503)
            update.assert_not_called()
        with patch.object(app, "update_templates") as update:
            for authorization in ("", "Bearer incorrect", "Basic test-admin-token", "Bearer test-admin-token extra"):
                self.assert_preserved(self.post(authorization=authorization), status=401)
            update.assert_not_called()

    def test_secrets_never_returned_or_logged_on_activation_error(self):
        secret = app.app.config["TEMPLATE_ADMIN_TOKEN"]
        with patch.object(app, "update_templates", side_effect=OSError(f"internal error {secret}")), self.assertLogs(app.app.logger, level="ERROR") as logs:
            response = self.post()
        self.assertEqual(500, response.status_code)
        self.assertNotIn(secret, response.get_data(as_text=True))
        self.assertNotIn(secret, " ".join(logs.output))
        for authorization in ("", "Bearer incorrect"):
            self.assertNotIn(secret, self.post(authorization=authorization).get_data(as_text=True))
        response = self.client.get("/")
        self.assertNotIn(secret, response.get_data(as_text=True))
        response.close()

    def test_uploaded_active_backup_and_recovery_files_are_not_public(self):
        backup = self.active.parent / ".template-backups" / ("b" * 32)
        backup.mkdir(parents=True)
        (backup / "template_en.docx").write_bytes(b"private backup")
        staged = self.active.parent / ".template-stage-private"
        staged.mkdir()
        (staged / "official_en.docx").write_bytes(b"private upload")
        for path in ("/templates/template_en.docx", "/backend/templates/template_en.docx",
                     "/.template-backups/" + "b" * 32 + "/template_en.docx",
                     "/.template-stage-private/official_en.docx", "/.template-transaction.json",
                     "/../backend/templates/template_en.docx", "/%2e%2e/backend/templates/template_en.docx"):
            with self.subTest(path=path):
                self.assertEqual(404, self.client.get(path).status_code)

    def test_failed_validation_cleans_staging_and_atomic_temporary_files(self):
        self.assert_preserved(self.post({"en": (io.BytesIO(b"invalid"), "official.docx")}))
        self.assertFalse(list(self.active.parent.glob(".template-stage-*")))
        self.assertFalse(list(self.active.parent.rglob(".template-copy-*")))
        self.assertFalse(list(self.active.parent.glob(".template-journal-*")))

    def test_macro_content_and_external_linked_content_rejected(self):
        for payload in (
            mutate_docx(self.sources["en"], lambda name, data: data.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.document.macroEnabled.main+xml") if name == "[Content_Types].xml" else data),
            mutate_docx(self.sources["en"], lambda name, data: data.replace(b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject") if name.endswith(".rels") else data),
        ):
            self.assert_preserved(self.post({"en": (io.BytesIO(payload), "official.docx")}))

    def test_corrupt_zip_crc_rejected(self):
        # ZIP_STORED makes the payload mutation reliably trigger a CRC failure.
        target = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.sources["en"])) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_STORED) as archive:
            for info in source.infolist():
                archive.writestr(info.filename, source.read(info.filename))
        data = target.getvalue().replace(b"[Name of academic unit]", b"[Name of academic uniu]", 1)
        self.assert_preserved(self.post({"en": (io.BytesIO(data), "official.docx")}))

    def test_interrupted_transaction_recovered_before_reader(self):
        backup_id = "a" * 32
        backup = self.active.parent / ".template-backups" / backup_id
        backup.mkdir(parents=True)
        for name, content in self.before.items():
            (backup / name).write_bytes(content)
        journal = self.active.parent / ".template-transaction.json"
        journal.write_text(json.dumps({"backup_id": backup_id}))
        (self.active / "template_en.docx").write_bytes(b"interrupted update")
        with template_storage.template_lock(self.active):
            self.assertEqual(self.before, self.snapshot())
        self.assertFalse(journal.exists())

    def test_activation_waits_for_generation_lock(self):
        started, finished = threading.Event(), threading.Event()
        errors = []
        def updater():
            started.set()
            try:
                template_storage.activate_templates(self.active, self.active)
            except Exception as exc:
                errors.append(exc)
            finally:
                finished.set()
        with template_storage.template_lock(self.active):
            thread = threading.Thread(target=updater)
            thread.start()
            self.assertTrue(started.wait(2))
            self.assertFalse(finished.wait(0.1))
        thread.join(5)
        self.assertTrue(finished.is_set())
        self.assertEqual([], errors)
        self.assertEqual(self.before, self.snapshot())


if __name__ == "__main__":
    unittest.main()
